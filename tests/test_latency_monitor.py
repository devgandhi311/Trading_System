"""Tests for execution latency monitoring."""

from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from execution import LatencyMonitor, LatencyThresholds, LATENCY_STAGES
from models import StrategyType


def test_latency_monitor_records_samples_and_classifies() -> None:
    monitor = LatencyMonitor()
    sample = monitor.record("ml_prediction", 750.0, strategy_type=StrategyType.CREDIT_SPREAD)

    assert "ml_prediction" in LATENCY_STAGES
    assert sample.latency_ms == 750.0
    assert monitor.latest_latency_ms() == 750.0
    assert monitor.classify_latency(100.0) == "paper_backtest_expected"
    assert monitor.classify_latency(3_000.0) == "normal_stable_setup"
    assert monitor.classify_latency(20_000.0) == "severe_latency"


def test_latency_policy_rejects_scalping_and_reduces_confidence() -> None:
    monitor = LatencyMonitor(thresholds=LatencyThresholds(scalping_reject_ms=2_000.0))
    decision = monitor.assess_trade_latency(
        strategy_type=StrategyType.VOLATILITY_SCALPING,
        confidence_score=0.8,
        latency_ms=3_000.0,
    )

    assert decision.allowed is False
    assert decision.reject_scalping is True
    assert decision.adjusted_confidence < 0.8
    assert decision.warning is not None


def test_latency_policy_allows_non_scalping_with_confidence_adjustment() -> None:
    monitor = LatencyMonitor()
    decision = monitor.assess_trade_latency(
        strategy_type=StrategyType.CREDIT_SPREAD,
        confidence_score=0.8,
        latency_ms=1_500.0,
    )

    assert decision.allowed is True
    assert decision.reject_scalping is False
    assert 0.0 < decision.adjusted_confidence < 0.8


def test_latency_monitor_stage_duration_and_total() -> None:
    monitor = LatencyMonitor()
    monitor._marks["decision_start"] = 100.0
    monitor._marks["exchange_acknowledgement"] = 101.25

    sample = monitor.record_total_decision_to_order(strategy_type="credit_spread")

    assert sample.component == "total_decision_to_order_time"
    assert sample.latency_ms == 1_250.0


def test_latency_monitor_saves_daily_report(tmp_path) -> None:
    monitor = LatencyMonitor(report_dir=tmp_path)
    monitor.record("live_tick_received", 20.0)
    monitor.record("risk_validation", 120.0)
    monitor.record("risk_validation", 180.0)

    report_path = monitor.save_daily_report(report_date=date.today())

    assert report_path.exists()
    workbook = load_workbook(report_path)
    sheet = workbook["Latency"]
    headers = [cell.value for cell in sheet[1]]
    components = [row[headers.index("component")].value for row in sheet.iter_rows(min_row=2)]
    assert "risk_validation" in components
    assert "sample:risk_validation" in components
